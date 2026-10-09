"""Reads the Secure Controls Framework workbook that the SCF Council publishes.

The workbook is read as downloaded. Nothing from it is stored in this repository: the
importer runs on the installer's own copy and writes only to their database.
"""

from __future__ import annotations

import re
from pathlib import Path

from .spec import FrameworkSpec, ImportProblem, NodeSpec

SHEET = re.compile(r"^SCF (\d{4}\.\d+)$")
CONTROL_ID = re.compile(r"^[A-Z]{2,4}-\d{2}(\.\d+)*$")
NOTE = (
    "Imported from a Secure Controls Framework workbook downloaded by the person who installed "
    "this system. The SCF Council publishes it under Creative Commons licensing; its terms "
    "(securecontrolsframework.com) say modified copies may not be redistributed, so this copy "
    "stays inside this installation and cannot be exported."
)
ATTRIBUTION = "Secure Controls Framework (SCF), SCF Council. https://securecontrolsframework.com"


def _clean(value) -> str:
    return re.sub(r"[ \t]+", " ", str(value or "").replace("\r", "")).strip()


def _headers(row) -> dict[str, int]:
    return {" ".join(str(h).split()): i for i, h in enumerate(row) if h}


def _tokens(cell) -> list[str]:
    return [t.strip() for t in re.split(r"[\n;,]+", str(cell or "")) if t.strip()]


def _iso_refs(clauses, controls) -> list[str]:
    refs = []
    for token in _tokens(clauses):
        ref = re.sub(r"\(.*$", "", token)  # 5.1(a) maps to clause 5.1
        if re.fullmatch(r"\d{1,2}(\.\d{1,2}){0,2}", ref):
            refs.append(ref)
    for token in _tokens(controls):
        if re.fullmatch(r"[5-8]\.\d{1,2}", token):  # ISO 27002 clauses 5 to 8 are Annex A
            refs.append(f"A.{token}")
    return sorted(set(refs))


def parse(path: str | Path) -> FrameworkSpec:
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ImportProblem(f"The file could not be opened as a workbook: {exc}") from None
    try:
        sheet_name = next((n for n in workbook.sheetnames if SHEET.match(n)), None)
        if sheet_name is None:
            raise ImportProblem(
                "No sheet named like 'SCF 2026.3' was found. Is this the SCF workbook?"
            )
        version = SHEET.match(sheet_name).group(1)
        rows = list(workbook[sheet_name].iter_rows(values_only=True))
        domain_names = _domain_names(workbook)
    finally:
        workbook.close()

    if not rows:
        raise ImportProblem("The SCF sheet is empty.")
    columns = _headers(rows[0])
    description_col = next(
        (i for h, i in columns.items() if h.startswith("Secure Controls Framework (SCF) Control")),
        None,
    )
    needed = {
        "SCF Domain": columns.get("SCF Domain"),
        "SCF Control": columns.get("SCF Control"),
        "SCF #": columns.get("SCF #"),
        "description": description_col,
    }
    missing = [k for k, v in needed.items() if v is None]
    if missing:
        raise ImportProblem(f"Expected columns are missing: {', '.join(missing)}.")
    question = columns.get("SCF Control Question")
    weighting = columns.get("Relative Control Weighting")
    function = columns.get("NIST CSF Function Grouping")
    iso_clauses, iso_controls = columns.get("ISO 27001 2022"), columns.get("ISO 27002 2022")
    csf = columns.get("NIST CSF 2.0")

    nodes: list[NodeSpec] = []
    known: set[str] = set()
    for row in rows[1:]:
        ref = _clean(row[needed["SCF #"]])
        if not ref:
            continue
        if not CONTROL_ID.match(ref):
            raise ImportProblem(f"Unexpected control number {ref!r}.")
        prefix = ref.split("-")[0]
        if prefix not in known:
            known.add(prefix)
            nodes.append(
                NodeSpec(
                    ref_id=prefix,
                    name=domain_names.get(prefix) or _clean(row[needed["SCF Domain"]]),
                    assessable=False,
                )
            )
        maps_to = {}
        if iso_clauses is not None or iso_controls is not None:
            refs = _iso_refs(
                row[iso_clauses] if iso_clauses is not None else None,
                row[iso_controls] if iso_controls is not None else None,
            )
            if refs:
                maps_to["iso-27001-2022"] = refs
        if csf is not None and _tokens(row[csf]):
            maps_to["nist-csf-2.0"] = _tokens(row[csf])
        weight = row[weighting] if weighting is not None else 1
        try:
            weight = max(1, min(10, int(float(weight))))
        except (TypeError, ValueError):
            weight = 1
        parent = prefix
        probe = ref
        while "." in probe:
            probe = probe.rsplit(".", 1)[0]
            if probe in known:
                parent = probe
                break
        known.add(ref)
        extra = {"maps_to": maps_to} if maps_to else {}
        if question is not None and _clean(row[question]):
            extra["question"] = _clean(row[question])
        if function is not None and _clean(row[function]):
            extra["csf_function"] = _clean(row[function])
        nodes.append(
            NodeSpec(
                ref_id=ref,
                name=_clean(row[needed["SCF Control"]]),
                description=_clean(row[needed["description"]]),
                parent_ref=parent,
                weight=weight,
                extra=extra,
            )
        )
    if sum(1 for n in nodes if n.assessable) < 100:
        raise ImportProblem(
            "Far fewer controls than expected were found. Is this the SCF workbook?"
        )
    spec = FrameworkSpec(
        slug=f"scf-{version}",
        name="Secure Controls Framework (SCF)",
        version=version,
        provider="SCF Council",
        nodes=nodes,
        description="A framework of frameworks with controls mapped to laws, regulations and standards.",
        licence_note=NOTE,
        attribution=ATTRIBUTION,
        source=f"Workbook {Path(path).name}, sheet {sheet_name}",
        redistributable=False,
    )
    return spec


def _domain_names(workbook) -> dict[str, str]:
    name = "SCF Domains & Principles"
    if name not in workbook.sheetnames:
        return {}
    rows = list(workbook[name].iter_rows(values_only=True))
    if not rows:
        return {}
    columns = _headers(rows[0])
    ident, label = columns.get("SCF Identifier"), columns.get("SCF Domain")
    if ident is None or label is None:
        return {}
    return {_clean(r[ident]): _clean(r[label]) for r in rows[1:] if r[ident] and r[label]}
